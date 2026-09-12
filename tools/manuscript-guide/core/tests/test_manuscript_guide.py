import argparse
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path


MODULE_PATH = Path(__file__).parents[1] / "manuscript_guide.py"
SPEC = importlib.util.spec_from_file_location("manuscript_guide", MODULE_PATH)
guide = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
SPEC.loader.exec_module(guide)


class ManuscriptGuideTests(unittest.TestCase):
    def test_locked_edit_survives_rebuild(self):
        generated = [{
            "id": "entity-abc", "canonical_name": "Arelian", "category": "Character",
            "pronunciation": {"say_as": "Arelian"}, "description": {"text": ""},
            "personality_notes": [], "locked_fields": [],
        }]
        previous = {"entities": [{
            "id": "entity-abc", "canonical_name": "Arelian", "category": "Character",
            "pronunciation": {"say_as": "ah-RELL-ee-in"}, "description": {"text": ""},
            "personality_notes": [], "locked_fields": ["pronunciation"],
        }]}
        merged = guide.merge_locked(generated, previous)
        self.assertEqual("ah-RELL-ee-in", merged[0]["pronunciation"]["say_as"])

    def test_direct_trait_has_evidence(self):
        occurrences = [{"chapter": "Chapter 1", "text": "Arelian was brave and wary.", "start": "0", "end": "7"}]
        notes = guide.trait_notes("Arelian", occurrences)
        self.assertEqual({"Described as brave.", "Described as wary."}, {note["text"] for note in notes})
        self.assertEqual("Chapter 1", notes[0]["evidence"]["chapter"])

    def test_fiction_fixture_extracts_candidates_and_source_evidence(self):
        paragraphs = [
            {"chapter": "Chapter 1", "text": "Captain Arelian said the Council of Ash would meet in Dawnspire."},
            {"chapter": "Chapter 1", "text": "Arelian was a veteran navigator, brave and wary before the council arrived."},
        ]
        entities = guide.build_entities(paragraphs, "model-that-is-not-installed", None)
        by_name = {entity["canonical_name"]: entity for entity in entities}
        self.assertIn("Captain Arelian", by_name)
        self.assertEqual("Character", by_name["Captain Arelian"]["category"])
        self.assertTrue(by_name["Captain Arelian"]["occurrences"][0]["excerpt"])
        self.assertEqual("Explicitly described as veteran navigator.", by_name["Captain Arelian"]["description"]["text"])

    def test_status_and_hotword_export_are_independent_files(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            manuscript = root / "Manuscript.docx"
            manuscript.write_bytes(b"fixture manuscript")
            guide_file = root / "ManuscriptGuide" / "manuscript_guide.json"
            guide.write_json(str(guide_file), {
                "source": {"sha256": guide.document_hash(str(manuscript))},
                "entities": [{"id": "entity-1", "canonical_name": "Dawnspire", "aliases": ["the Spire"]}],
            })
            status_file = root / "ManuscriptGuide" / "status.txt"
            guide.status(argparse.Namespace(docx=str(manuscript), guide=str(guide_file), out=str(status_file)))
            self.assertEqual("STATUS|CURRENT\n", status_file.read_text(encoding="utf-8"))
            hotwords = root / "ManuscriptGuide" / "whisper_hotwords.txt"
            guide.export_hotwords(argparse.Namespace(guide=str(guide_file), out=str(hotwords), entity_ids=""))
            self.assertIn("Dawnspire", hotwords.read_text(encoding="utf-8"))
            self.assertFalse((root / "TranscriptCompare").exists())

    def test_real_docx_build_uses_project_owned_guide_file(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            manuscript = root / "Manuscript.docx"
            document = Document()
            document.add_heading("Chapter 1", level=1)
            document.add_paragraph("Captain Arelian said the Council of Ash would meet in Dawnspire.")
            document.save(manuscript)
            output = root / "ManuscriptGuide" / "manuscript_guide.json"
            guide.build(argparse.Namespace(
                docx=str(manuscript), out=str(output), progress=str(root / "ManuscriptGuide" / "progress.txt"),
                spacy_model="en_core_web_sm", espeak_library="",
            ))
            data = json.loads(output.read_text(encoding="utf-8"))
            self.assertEqual(guide.document_hash(str(manuscript)), data["source"]["sha256"])
            self.assertTrue(data["entities"])
            self.assertTrue((root / "ManuscriptGuide" / "progress.txt").read_text(encoding="utf-8").startswith("DONE|100"))


if __name__ == "__main__":
    unittest.main()
