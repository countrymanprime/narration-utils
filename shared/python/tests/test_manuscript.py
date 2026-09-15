import sys
import types

import pytest
from docx import Document

from narration_common import manuscript


def test_docx_import_splits_chapters_and_skips_table_of_contents(tmp_path):
    doc = Document()
    doc.add_paragraph("Table of Contents", style="Heading 1")
    doc.add_paragraph("Should not appear as its own chapter.")
    doc.add_paragraph("Chapter One", style="Heading 1")
    doc.add_paragraph("First   paragraph\twith odd whitespace.")
    doc.add_paragraph("")  # empty paragraphs are dropped
    source = tmp_path / "book.docx"
    doc.save(source)

    draft = manuscript.prepare_import(source)
    assert draft["chapterTitles"] == ["Chapter One"]
    assert [item["chapter"] for item in draft["paragraphs"]] == ["Front Matter", "Chapter One"]
    assert draft["paragraphs"][1]["text"] == "First paragraph with odd whitespace."

    committed = manuscript.commit_import(tmp_path, source, draft)
    assert [c["title"] for c in committed["chapters"]] == ["Front Matter", "Chapter One"]


def test_markdown_import_uses_selected_chapter_level_and_keeps_deeper_sections(tmp_path):
    source = tmp_path / "book.md"
    source.write_text("Preface text.\n\n# Chapter One\n\nOpening text.\n\n## A section\n\nSection text.\n", encoding="utf-8")

    draft = manuscript.prepare_import(source, markdown_heading_level=1)
    assert [item["chapter"] for item in draft["paragraphs"]] == ["Front Matter", "Chapter One", "Chapter One"]
    assert draft["paragraphs"][-1]["section"] == "A section"

    committed = manuscript.commit_import(tmp_path, source, draft)
    loaded = manuscript.load(tmp_path)
    assert loaded["documentId"] == committed["documentId"]
    assert loaded["paragraphs"][1]["id"] == "p-000002"
    assert (tmp_path / loaded["source"]["storedPath"]).read_text(encoding="utf-8") == source.read_text(encoding="utf-8")


def test_import_splits_a_short_cover_page_from_remaining_front_matter(tmp_path):
    source = tmp_path / "book.md"
    source.write_text("The Lantern\n\nBy Example Author\n\nCopyright 2026\n\nA preface for readers.\n\n# Chapter One\n\nOpening text.", encoding="utf-8")

    draft = manuscript._markdown_draft(source, 1)

    assert [item["chapter"] for item in draft["paragraphs"]] == ["Cover", "Cover", "Cover", "Front Matter", "Chapter One"]


def test_docx_heading_line_break_becomes_a_chapter_subtitle(tmp_path):
    doc = Document()
    heading = doc.add_heading("Chapter One", level=1)
    heading.add_run().add_break()
    heading.add_run("Bad Ideas")
    doc.add_paragraph("Opening text.")
    source = tmp_path / "book.docx"
    doc.save(source)

    draft = manuscript._docx_draft(source)
    committed = manuscript.commit_import(tmp_path, source, draft)

    assert draft["chapterTitles"] == ["Chapter One"]
    assert committed["chapters"][0]["subtitle"] == "Bad Ideas"


def test_failed_import_does_not_replace_active_manuscript(tmp_path):
    source = tmp_path / "book.md"
    source.write_text("# One\n\nFirst text.", encoding="utf-8")
    manuscript.commit_import(tmp_path, source, manuscript.prepare_import(source))
    before = manuscript.manuscript_path(tmp_path).read_text(encoding="utf-8")

    missing = tmp_path / "missing.md"
    with pytest.raises(manuscript.ManuscriptError):
        manuscript.prepare_import(missing)

    assert manuscript.manuscript_path(tmp_path).read_text(encoding="utf-8") == before


def test_scanned_pdf_is_rejected(monkeypatch, tmp_path):
    source = tmp_path / "scan.pdf"
    source.write_bytes(b"not really a pdf")

    class Reader:
        def __init__(self, _path):
            self.pages = [types.SimpleNamespace(extract_text=lambda: "")]

    monkeypatch.setitem(sys.modules, "pypdf", types.SimpleNamespace(PdfReader=Reader))
    with pytest.raises(manuscript.ManuscriptError, match="selectable"):
        manuscript.prepare_import(source)
